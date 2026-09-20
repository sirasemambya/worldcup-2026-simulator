#!/usr/bin/env python3
"""
report.py — Generate formatted Excel report from model_probs.csv
Usage: python3 report.py
"""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from datetime import datetime

INPUT  = "model_probs.csv"
OUTPUT = f"wc2026_report_{datetime.now().strftime('%Y-%m-%d_%H%M%S')}.xlsx"

# ── Color palette ─────────────────────────────────────────────────────────────
NAVY        = "0D1B2A"
GOLD        = "C9A227"
GOLD_LIGHT  = "FFF8DC"
GOLD_MED    = "FFE680"
GREEN_LIGHT = "D6F0E0"
BLUE_LIGHT  = "DDEEFF"
GRAY_LIGHT  = "F5F6F8"
GRAY_MED    = "E2E5EA"
WHITE       = "FFFFFF"

GROUP_HEADER_FILLS = [
    "1B3A5C", "1A4A3A", "3A1A4A", "4A2A1A",
    "1A3A4A", "3A3A1A", "1A1A4A", "2A3A2A",
    "4A1A2A", "2A4A1A", "3A1A3A", "1A4A1A",
]


def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)


def font(bold=False, color="111111", size=11):
    return Font(bold=bold, color=color, name="Calibri", size=size)


def border():
    s = Side(style="thin", color="D0D5DD")
    return Border(left=s, right=s, top=s, bottom=s)


def center_align():
    return Alignment(horizontal="center", vertical="center")


def left_align(indent=0):
    return Alignment(horizontal="left", vertical="center", indent=indent)


def set_col_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_title(ws, text, n_cols):
    """Full-width title banner at the top of each sheet."""
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    cell = ws.cell(row=1, column=1, value=text)
    cell.fill = fill(NAVY)
    cell.font = Font(bold=True, color=GOLD, name="Calibri", size=14)
    cell.alignment = center_align()
    ws.row_dimensions[1].height = 32


def write_col_headers(ws, row, labels):
    for col, label in enumerate(labels, 1):
        cell = ws.cell(row=row, column=col, value=label)
        cell.fill = fill("1E3A5F")
        cell.font = font(bold=True, color="FFFFFF", size=10)
        cell.alignment = center_align()
        cell.border = border()
    ws.row_dimensions[row].height = 20


# ── Group stage sheets ────────────────────────────────────────────────────────

def write_group_sheet(ws, df, prob_col, title):
    ws.sheet_view.showGridLines = False
    COLS   = ["Team", "Probability"]
    WIDTHS = [26, 16]
    set_col_widths(ws, WIDTHS)

    write_title(ws, title, 2)
    write_col_headers(ws, 2, COLS)

    row = 3
    for g_idx, grp in enumerate(sorted(df["group"].unique())):
        gdf = df[df["group"] == grp].sort_values(prob_col, ascending=False)

        # Group header
        header_color = GROUP_HEADER_FILLS[g_idx % len(GROUP_HEADER_FILLS)]
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        cell = ws.cell(row=row, column=1, value=f"  GROUP  {grp}")
        cell.fill = fill(header_color)
        cell.font = Font(bold=True, color=GOLD, name="Calibri", size=11)
        cell.alignment = left_align(indent=1)
        cell.border = border()
        ws.row_dimensions[row].height = 22
        row += 1

        for rank, (_, r) in enumerate(gdf.iterrows()):
            prob = r[prob_col]
            if rank == 0:
                row_fill = GOLD_LIGHT
                bold = True
            elif rank == 1:
                row_fill = GREEN_LIGHT
                bold = False
            elif rank == 2:
                row_fill = BLUE_LIGHT
                bold = False
            else:
                row_fill = GRAY_LIGHT
                bold = False

            # Team name
            c1 = ws.cell(row=row, column=1, value=r["team"])
            c1.fill = fill(row_fill)
            c1.font = font(bold=bold, color="111111")
            c1.alignment = left_align(indent=3)
            c1.border = border()

            # Probability
            c2 = ws.cell(row=row, column=2, value=f"{prob:.1f}%")
            c2.fill = fill(row_fill)
            c2.font = font(bold=bold, color="111111")
            c2.alignment = center_align()
            c2.border = border()

            ws.row_dimensions[row].height = 20
            row += 1

        # Gap between groups
        for col in range(1, 3):
            cell = ws.cell(row=row, column=col, value="")
            cell.fill = fill(WHITE)
        ws.row_dimensions[row].height = 8
        row += 1


# ── Knockout round sheets ─────────────────────────────────────────────────────

def write_knockout_sheet(ws, df, prob_col, title):
    ws.sheet_view.showGridLines = False
    COLS   = ["#", "Team", "Group", "Probability"]
    WIDTHS = [6, 26, 10, 16]
    set_col_widths(ws, WIDTHS)

    write_title(ws, title, 4)
    write_col_headers(ws, 2, COLS)

    sdf = df[df[prob_col] > 0].sort_values(prob_col, ascending=False).reset_index(drop=True)

    for i, (_, r) in enumerate(sdf.iterrows()):
        prob = r[prob_col]
        row  = i + 3

        if i < 4:
            row_fill = GOLD_LIGHT
            bold = True
        elif i < 8:
            row_fill = GREEN_LIGHT
            bold = False
        elif i < 16:
            row_fill = BLUE_LIGHT
            bold = False
        elif i % 2 == 0:
            row_fill = GRAY_LIGHT
            bold = False
        else:
            row_fill = WHITE
            bold = False

        values = [i + 1, r["team"], r["group"], f"{prob:.1f}%"]
        aligns = [center_align(), left_align(indent=2), center_align(),
                  center_align()]

        for col, (val, aln) in enumerate(zip(values, aligns), 1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.fill = fill(row_fill)
            cell.font = font(bold=bold, color="111111")
            cell.alignment = aln
            cell.border = border()

        ws.row_dimensions[row].height = 20


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    df = pd.read_csv(INPUT)

    wb = Workbook()
    wb.remove(wb.active)

    # Group stage sheets
    ws1 = wb.create_sheet("Group Win")
    ws1.sheet_properties.tabColor = "C9A227"
    write_group_sheet(ws1, df, "group_win", "2026 World Cup — Group Win Probability")

    ws2 = wb.create_sheet("Advance (Top 32)")
    ws2.sheet_properties.tabColor = "2E86AB"
    write_group_sheet(ws2, df, "advance", "2026 World Cup — Advance from Group (Top 32)")

    ws3 = wb.create_sheet("Finish 2nd")
    ws3.sheet_properties.tabColor = "E05C00"
    write_group_sheet(ws3, df, "group_2nd", "2026 World Cup — Finish 2nd in Group")

    ws4 = wb.create_sheet("Finish Last")
    ws4.sheet_properties.tabColor = "CC2200"
    write_group_sheet(ws4, df, "group_last", "2026 World Cup — Finish Last in Group (4th)")

    # Knockout sheets
    ko = [
        ("Round of 16",   "R16",    "3A86FF"),
        ("Quarter-Final", "QF",     "8338EC"),
        ("Semi-Final",    "SF",     "FB5607"),
        ("Final",         "Final",  "FF006E"),
        ("Winner",        "Winner", "1A6B3C"),
    ]

    for sheet_name, col, tab_color in ko:
        ws = wb.create_sheet(sheet_name)
        ws.sheet_properties.tabColor = tab_color
        write_knockout_sheet(ws, df, col, f"2026 World Cup — {sheet_name}")

    wb.save(OUTPUT)
    print(f"Saved {OUTPUT}")


if __name__ == "__main__":
    main()
